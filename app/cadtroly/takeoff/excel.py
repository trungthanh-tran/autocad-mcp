"""Xuất Excel có truy vết: sheet Khối lượng, Chi tiết (từng handle), Ánh xạ, Kiểm soát."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..reader.model import Drawing
from .calc import TakeoffRow
from .mapper import Group

FILL = {"A": "C6EFCE", "B": "FFEB9C", "C": "FFC7CE"}
# Căn cứ ghi vào sheet Kiểm soát. Mã định mức do kỹ sư gán (takeoff_decide / tham số ma_dinh_muc), tool không tự điền.
CAN_CU = ("Nghị định 206/2026/NĐ-CP (quản lý chi phí đầu tư xây dựng, hiệu lực 01/7/2026); "
          "Thông tư 38/2026/TT-BXD (định mức xây dựng, hiệu lực 01/7/2026, thay TT 12/2021/TT-BXD)")


def _autosize(ws, maxw=60):
    for col in ws.columns:
        w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, w + 2), maxw)


def export_excel(dwg: Drawing, groups: list[Group], rows: list[TakeoffRow], out: str | Path,
                 *, meta: dict | None = None) -> Path:
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    bold = Font(bold=True)

    # --- Khối lượng ---
    ws = wb.active
    ws.title = "Khoi luong"
    ws.append(["BẢNG KHỐI LƯỢNG", "", "", "", "", "", ""])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Bản vẽ", dwg.path])
    ws.append(["Đơn vị bản vẽ", dwg.units, "Xuất lúc", datetime.now().strftime("%d/%m/%Y %H:%M")])
    ws.append([])
    hdr = ["STT", "Mã định mức", "Công tác", "ĐV", "Khối lượng", "Cách tính", "Nguồn (truy vết)", "Tin cậy", "Nhóm", "Ghi chú"]
    ws.append(hdr)
    for c in ws[ws.max_row]:
        c.font = bold
    for i, r in enumerate(rows, 1):
        ws.append([i, r.ma_dinh_muc or None, r.cong_tac, r.don_vi, r.khoi_luong, r.cach_tinh, r.nguon, r.tin_cay,
                   r.nhom, r.ghi_chu])
        cell = ws.cell(row=ws.max_row, column=8)
        cell.fill = PatternFill("solid", fgColor=FILL.get(r.tin_cay, "FFFFFF"))
        cell.alignment = Alignment(horizontal="center")
    ws.append([])
    ws.append(["A = đếm/đo trực tiếp theo quy tắc", "B = suy từ tên hoặc text, cần xác nhận", "C = không xác định, để trống"])
    _autosize(ws)

    # --- Chi tiết ---
    wd = wb.create_sheet("Chi tiet")
    wd.append(["Công tác", "Handle", "Layer/Block", "Giá trị", "Đơn vị", "X", "Y", "Độ chính xác", "Lệnh zoom"])
    for c in wd[1]:
        c.font = bold
    for r in rows:
        for d in r.chi_tiet:
            p = d.get("vi_tri") or (None, None)
            zoom = f"_.ZOOM _C {p[0]:.0f},{p[1]:.0f} 5000" if p and p[0] is not None else ""
            wd.append([r.cong_tac, d["handle"], r.nhom, d.get("gia_tri"), r.don_vi,
                       round(p[0], 1) if p[0] is not None else None,
                       round(p[1], 1) if p[1] is not None else None,
                       d.get("do_chinh_xac", "exact"), zoom])
    _autosize(wd)

    # --- Ánh xạ ---
    wm = wb.create_sheet("Anh xa")
    wm.append(["Loại", "Tên trong bản vẽ", "Số lượng", "Mô tả", "Công tác", "ĐV", "Cách tính", "Tham số", "Tin cậy", "Lý do"])
    for c in wm[1]:
        c.font = bold
    for g in groups:
        wm.append([g.kind, g.name, g.count, g.describe(), "(bỏ qua)" if g.skip else g.cong_tac, g.don_vi,
                   g.cach_tinh, str(g.tham_so) if g.tham_so else "", g.tin_cay, g.ly_do])
        wm.cell(row=wm.max_row, column=9).fill = PatternFill("solid", fgColor=FILL.get(g.tin_cay, "FFFFFF"))
    _autosize(wm)

    # --- Kiểm soát ---
    wk = wb.create_sheet("Kiem soat")
    wk.append(["Mục", "Giá trị"])
    for c in wk[1]:
        c.font = bold
    s = dwg.summary()
    for k, v in [("Đối tượng", s["entities"]), ("Layer", s["layers"]), ("Block", s["blocks"]),
                 ("INSERT", s["inserts"]), ("Loại", str(s["kinds"]))]:
        wk.append([k, v])
    for k, v in (meta or {}).items():
        wk.append([k, str(v)])
    wk.append(["Phương pháp", "Đếm block theo tên; đo polyline/hatch flatten dung sai 0.01; cao từ text H= hoặc người nhập; không bịa số thiếu."])
    _control_tail(wk, rows=[r.ma_dinh_muc for r in rows])
    _autosize(wk)

    wb.save(out)
    return out


def _control_tail(wk, *, rows: list[str]) -> None:
    wk.append(["Căn cứ", CAN_CU])
    thieu = sum(1 for m in rows if not m)
    wk.append(["Mã định mức", f"{len(rows) - thieu}/{len(rows)} dòng đã gán" +
               (" — dòng trống cần kỹ sư gán theo TT 38/2026/TT-BXD" if thieu else "")])


def _header(ws, cols: list[str]) -> None:
    ws.append(cols)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)


def _fill(ws, col: int, tc: str) -> None:
    c = ws.cell(row=ws.max_row, column=col)
    c.fill = PatternFill("solid", fgColor=FILL.get(tc, "FFFFFF"))
    c.alignment = Alignment(horizontal="center")


def export_earthwork(dwg: Drawing, res: dict, out: str | Path, *, ma_dinh_muc: dict | None = None) -> Path:
    """Excel đào đắp: Tổng hợp, Khối lượng từng đoạn, Trắc ngang (diện tích + handle), Kiểm soát."""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    ma = ma_dinh_muc or {}
    names, loai = res["names"], sorted(res["totals"])
    wb = Workbook()

    ws = wb.active
    ws.title = "Tong hop"
    ws.append(["KHỐI LƯỢNG ĐÀO ĐẮP (mặt cắt trung bình)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Bản vẽ", dwg.path])
    ws.append(["Số trắc ngang", len(res["sections"]), "Xuất lúc", datetime.now().strftime("%d/%m/%Y %H:%M")])
    ws.append([])
    _header(ws, ["STT", "Mã định mức", "Công tác", "ĐV", "Khối lượng", "Tin cậy", "Ghi chú"])
    for i, k in enumerate(loai, 1):
        tcs = [s["tin_cay"].get(k, "A") for s in res["segments"] if k in s["tin_cay"]]
        tc = max(tcs) if tcs else "C"
        ws.append([i, ma.get(k) or None, names.get(k, k), "m3", res["totals"][k], tc,
                   "có đoạn một đầu không ghi diện tích, coi = 0" if tc == "B" else ""])
        _fill(ws, 6, tc)
    for w in res["warnings"]:
        ws.append(["Cảnh báo", w])
    _autosize(ws)

    wseg = wb.create_sheet("Khoi luong doan")
    _header(wseg, ["Từ", "Đến", "L (m)"] + [f"V {names.get(k, k)} (m3)" for k in loai])
    for s in res["segments"]:
        wseg.append([s["tu"], s["den"], s["L_m"]] + [s["khoi_luong"].get(k) for k in loai])
    _autosize(wseg)

    wsec = wb.create_sheet("Trac ngang")
    _header(wsec, ["Lý trình", "Lý trình (m)", "Handle lý trình"] + [f"F {names.get(k, k)} (m2)" for k in loai] +
            ["Handle diện tích", "Tin cậy", "Lệnh zoom", "Ghi chú"])
    for s in res["sections"]:
        tc = max(s.tin_cay.values()) if s.tin_cay else "C"
        hs = sorted({h for v in s.sources.values() for h in v})
        zoom = f"_.ZOOM _C {s.point[0]:.0f},{s.point[1]:.0f} 50" if s.point else ""
        wsec.append([s.label, s.station_m, s.handle] + [s.areas.get(k) for k in loai] +
                    [", ".join(hs), tc, zoom, "; ".join(s.ghi_chu) or ("không thấy diện tích" if not s.areas else "")])
        _fill(wsec, 4 + len(loai) + 1, tc)
    _autosize(wsec)

    wk = wb.create_sheet("Kiem soat")
    _header(wk, ["Mục", "Giá trị"])
    wk.append(["Phương pháp", "V = (F1 + F2) / 2 × L giữa hai trắc ngang liền kề, sắp theo lý trình. "
                              "Diện tích đọc từ text trên bản vẽ, không đo lại hình học."])
    wk.append(["Ghép", "Diện tích gán về text lý trình gần nhất; gần hai lý trình tương đương → tin cậy B."])
    _control_tail(wk, rows=[ma.get(k, "") for k in loai])
    _autosize(wk)
    wb.save(out)
    return out


def export_rebar(dwg: Drawing, items: list, by_d: dict, out: str | Path, *, ma_dinh_muc: dict | None = None) -> Path:
    """Excel thống kê thép: Thống kê (từng ghi chú), Tổng hợp theo đường kính, Kiểm soát."""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    ma = {int(k): v for k, v in (ma_dinh_muc or {}).items()}
    wb = Workbook()
    ws = wb.active
    ws.title = "Thong ke thep"
    ws.append(["BẢNG THỐNG KÊ CỐT THÉP (từ ghi chú trên bản vẽ)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Bản vẽ", dwg.path])
    ws.append([])
    _header(ws, ["STT", "Số hiệu", "Ø (mm)", "Số thanh", "Dài 1 thanh (m)", "Tổng dài (m)", "kg/m", "Tổng KL (kg)",
                 "Ghi chú gốc", "Handle", "Tin cậy", "Ghi chú"])
    for i, it in enumerate(items, 1):
        ws.append([i, it.so_hieu or None, it.d_mm, it.so_thanh, it.dai_m, it.tong_dai_m, it.kg_per_m, it.tong_kg,
                   it.text, ", ".join(it.handles), it.tin_cay, it.ghi_chu])
        _fill(ws, 11, it.tin_cay)
    _autosize(ws)

    wd = wb.create_sheet("Tong hop")
    _header(wd, ["Mã định mức", "Đường kính", "Số dòng", "Tổng dài (m)", "Tổng KL (kg)", "Dòng thiếu dữ liệu"])
    for d, g in by_d.items():
        wd.append([ma.get(d) or None, f"Ø{d}", g["so_dong"], g["tong_dai_m"], g["tong_kg"], g["thieu"]])
    wd.append(["", "Cộng", sum(g["so_dong"] for g in by_d.values()), None,
               round(sum(g["tong_kg"] for g in by_d.values()), 2), sum(g["thieu"] for g in by_d.values())])
    _autosize(wd)

    wk = wb.create_sheet("Kiem soat")
    _header(wk, ["Mục", "Giá trị"])
    wk.append(["Phương pháp", "KL = số thanh × dài × 0.0061654·d² kg/m (thép 7850 kg/m3). Đọc ghi chú dạng "
                              "'N1 12Ø16 L=5200', '%%c' = Ø; L < 50 coi là mét, còn lại mm."])
    wk.append(["Trùng", "Cùng số hiệu/Ø/số thanh/dài ở nhiều hình chiếu chỉ tính 1 lần (tin cậy B) — kiểm tra lại."])
    wk.append(["Thiếu", "Ghi chú chỉ có bước (a200) hoặc không có L= → tin cậy C, KL để trống."])
    _control_tail(wk, rows=[ma.get(d, "") for d in by_d])
    _autosize(wk)
    wb.save(out)
    return out
