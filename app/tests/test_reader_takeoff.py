import math
from pathlib import Path

from cadtroly.kho.rules import MappingRule, RuleStore, heuristic, norm
from cadtroly.reader import read_dxf
from cadtroly.takeoff.calc import compute, find_height, stats
from cadtroly.takeoff.excel import export_excel
from cadtroly.takeoff.mapper import apply_decision, propose_mapping


def test_read_summary(sample_dxf):
    d = read_dxf(sample_dxf)
    s = d.summary()
    assert d.units == "mm" and d.to_m == 0.001
    assert s["inserts"] == 13 and s["blocks"] == 3
    assert s["kinds"]["HATCH"] == 3 and s["kinds"]["DIMENSION"] == 2
    # hatch 3780 x 4840 mm = 18.295 m2
    h = [e for e in d.entities if e.dxftype == "HATCH"][0]
    assert math.isclose(h.area * 1e-6, 3.78 * 4.84, rel_tol=1e-3)
    # polyline tường ngoài kín, chu vi 40 m
    w = [e for e in d.entities if e.layer == "TUONG220"][0]
    assert math.isclose(w.length * 1e-3, 40.0, rel_tol=1e-6)
    # dimension override được nhận ra
    dims = [e for e in d.entities if e.dxftype == "DIMENSION"]
    assert sum(1 for e in dims if e.extra.get("override")) == 1


def test_norm_and_heuristic():
    assert norm("Tường-220") == norm("TUONG_220") == "tuong220"
    assert heuristic("layer", "TUONG220")[0] == "Xây tường 220"
    assert heuristic("layer", "TUONG220")[3]["day"] == 0.22
    assert heuristic("block", "CUA-D1")[0] == "Cửa đi"
    assert heuristic("layer", "A-GLAZ")[0] == "Cửa sổ"
    assert heuristic("layer", "DIM")[2] == "skip"
    assert heuristic("layer", "XYZ") is None


def test_propose_then_decide_then_rule_applies(sample_dxf, tmp_path):
    d = read_dxf(sample_dxf)
    rules = RuleStore(tmp_path / "quy-tac")
    groups = propose_mapping(d, rules)
    by = {(g.kind, g.name): g for g in groups}
    assert by[("block", "CUA-D1")].cong_tac.startswith("Cửa đi") and "900x2200" in by[("block", "CUA-D1")].cong_tac
    assert by[("block", "WD_2")].tin_cay == "B" and by[("block", "WD_2")].cong_tac == "Cửa sổ"
    assert by[("layer", "0")].tin_cay == "C"
    assert ("layer", "DIM") not in by and ("layer", "TEXT") not in by   # text/dim không thành nhóm bóc
    s0 = stats(groups)
    assert s0["A"] == 0
    apply_decision(by[("block", "WD_2")], rules, cong_tac="Cửa sổ nhôm 1200x1400", don_vi="bộ", cach_tinh="count")
    assert rules.path.exists() and rules.find("block", "wd_2").cong_tac == "Cửa sổ nhôm 1200x1400"
    groups2 = propose_mapping(d, RuleStore(tmp_path / "quy-tac"))
    g = [g for g in groups2 if g.name == "WD_2"][0]
    assert g.tin_cay == "A" and "don-vi" in g.ly_do
    assert stats(groups2)["A"] > 0


def test_compute_quantities(sample_dxf, tmp_path):
    d = read_dxf(sample_dxf)
    H, src = find_height(d)
    assert math.isclose(H, 3.3) and "H=3300" in src
    groups = propose_mapping(d, RuleStore(tmp_path / "q"))
    pending = {r.cong_tac: r for r in compute(d, groups)}
    # không suy diễn: nhóm B chưa được xác nhận → để trống, ghi lý do
    assert pending["Xây tường 220"].khoi_luong is None and pending["Xây tường 220"].cho_xac_nhan
    assert any("chiều cao 3.3 m lấy từ text" in x for x in pending["Xây tường 220"].cho_xac_nhan)
    from cadtroly.takeoff.calc import takeoff_questions
    assert {q["ma"] for q in takeoff_questions(list(pending.values()))} == {"anh_xa", "chieu_cao"}
    rows = {r.cong_tac: r for r in compute(d, groups, chap_nhan_B=True)}
    assert math.isclose(rows["Xây tường 220"].khoi_luong, 40 * 3.3 * 0.22, rel_tol=1e-3)
    assert rows["Xây tường 220"].tin_cay == "B"          # cao suy từ text → B
    assert math.isclose(rows["Xây tường 100"].khoi_luong, (5 + 5 + 12) * 3.3 * 0.1, rel_tol=1e-3)
    assert math.isclose(rows["Sàn BTCT"].khoi_luong, 3 * 3.78 * 4.84, rel_tol=1e-3)
    assert rows["Cửa sổ"].khoi_luong == 6 and rows["Cột BTCT"].khoi_luong == 4
    c = [r for r in rows.values() if r.tin_cay == "C"]
    assert c and c[0].khoi_luong is None                 # không bịa số
    rows_h = {r.cong_tac: r for r in compute(d, groups, height_m=3.6, chap_nhan_B=True)}
    assert math.isclose(rows_h["Xây tường 220"].khoi_luong, 40 * 3.6 * 0.22, rel_tol=1e-3)
    assert rows_h["Xây tường 220"].tin_cay == "B"        # vẫn B vì ánh xạ là đoán


def test_missing_height_leaves_blank(tmp_path):
    import ezdxf
    doc = ezdxf.new(); doc.header["$INSUNITS"] = 4
    doc.layers.add("TUONG220")
    doc.modelspace().add_lwpolyline([(0, 0), (1000, 0)], dxfattribs={"layer": "TUONG220"})
    p = tmp_path / "t.dxf"; doc.saveas(p)
    d = read_dxf(p)
    rows = compute(d, propose_mapping(d, RuleStore(tmp_path / "q")))
    assert rows[0].khoi_luong is None and rows[0].tin_cay == "C" and "chiều cao" in rows[0].ghi_chu


def test_export_excel(sample_dxf, tmp_path):
    from openpyxl import load_workbook
    d = read_dxf(sample_dxf)
    groups = propose_mapping(d, RuleStore(tmp_path / "q"))
    rows = compute(d, groups)
    out = export_excel(d, groups, rows, tmp_path / "out" / "kl.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Khoi luong", "Chi tiet", "Anh xa", "Kiem soat"]
    ws = wb["Chi tiet"]
    handles = [r[1].value for r in ws.iter_rows(min_row=2) if r[1].value]
    assert len(handles) >= 13 and all(isinstance(h, str) for h in handles)
    assert any(str(r[8].value).startswith("_.ZOOM") for r in ws.iter_rows(min_row=2))
