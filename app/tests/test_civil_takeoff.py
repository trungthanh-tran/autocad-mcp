"""Bóc khối lượng cầu đường (đào đắp từ trắc ngang, thép từ ghi chú) và mã định mức."""
import json
import math

import pytest
from openpyxl import load_workbook

from cadtroly.reader import read_dxf
from cadtroly.samples import make_cross_sections, make_rebar_notes
from cadtroly.takeoff.civil import earthwork, parse_station, rebar, rebar_by_diameter


def _data(r):
    return r.data if getattr(r, "data", None) is not None else json.loads(r.content[0].text)


def test_parse_station():
    assert parse_station("Km0+100.50") == 100.5
    assert parse_station("KM 2 + 040") == 2040
    assert parse_station("LÝ TRÌNH: 1+250") == 1250
    assert parse_station("Cọc C12") is None


def test_earthwork_asks_when_end_missing(tmp_path):
    res = earthwork(read_dxf(make_cross_sections(tmp_path / "tn.dxf")))
    assert [q["ma"] for q in res["can_hoi"]] == ["thieu_dien_tich"]      # Km0+100 không ghi F đắp
    assert res["totals"]["dap"] == 60.0 + 180.0                           # đoạn cuối chưa tính
    assert res["segments"][-1]["khoi_luong"]["dap"] is None and "dap" in res["incomplete"]
    assert res["totals"]["dao"] == 690.0                                  # đào đủ dữ liệu, không hỏi


def test_earthwork(tmp_path):
    res = earthwork(read_dxf(make_cross_sections(tmp_path / "tn.dxf")), thieu_dien_tich="bang_0")
    assert res["can_hoi"] == []
    secs = {s.station_m: s for s in res["sections"]}
    assert sorted(secs) == [0, 20, 50, 100]
    assert secs[20].areas == {"dao": 14.0, "dap": 4.0}          # "FĐào=14", "4,0"
    assert secs[50].tin_cay == {"dao": "B", "dap": "B"}          # ghép nhãn với ô số
    # (10+14)/2·20 + (14+6)/2·30 + (6+0)/2·50
    assert res["totals"]["dao"] == 690.0
    # (2+4)/2·20 + (4+8)/2·30 + (8+0)/2·50 — người dùng xác nhận đầu không ghi = 0
    assert res["totals"]["dap"] == 440.0
    assert res["segments"][0]["tin_cay"] == {"dao": "A", "dap": "A"}


def test_rebar_asks(tmp_path):
    from cadtroly.takeoff.civil import rebar_questions
    items = rebar(read_dxf(make_rebar_notes(tmp_path / "ct.dxf")))
    assert {q["ma"] for q in rebar_questions(items)} == {"don_vi_L", "trung_lap"}
    n1 = next(i for i in items if i.so_hieu == "N1")
    assert n1.dai_m is None and n1.tong_kg is None                        # chưa rõ đơn vị, chưa rõ trùng
    n2 = next(i for i in items if i.so_hieu == "N2")
    assert n2.dai_m == 12.5 and n2.tong_kg is not None                   # 'L=12.5m' ghi rõ đơn vị


def test_rebar_each_occurrence(tmp_path):
    items = rebar(read_dxf(make_rebar_notes(tmp_path / "ct.dxf")), don_vi_L="mm", trung_lap="tung_cho")
    assert sum(1 for i in items if i.so_hieu == "N1") == 2               # mỗi chỗ một thanh riêng


def test_rebar(tmp_path):
    items = rebar(read_dxf(make_rebar_notes(tmp_path / "ct.dxf")), don_vi_L="mm", trung_lap="mot_lan")
    by = {(i.so_hieu, i.d_mm): i for i in items}
    n1 = by[("N1", 16)]
    assert n1.so_thanh == 12 and n1.dai_m == 5.2 and len(n1.handles) == 2 and n1.tin_cay == "A"   # trùng, đã trả lời
    assert math.isclose(n1.tong_kg, 12 * 5.2 * 1.58, abs_tol=0.01)          # bảng TCVN 1651-2: D16 = 1,58 kg/m
    assert by[("N2", 22)].dai_m == 12.5                           # 'L=12.5m' ghi đơn vị
    assert by[("3", 12)].so_thanh == 20                           # 2x10
    n4 = by[("N4", 10)]
    assert n4.so_thanh is None and n4.tin_cay == "C" and "200" in n4.ghi_chu
    assert by[("N5", 25)].tong_kg is None
    by_d = rebar_by_diameter(items)
    assert by_d[10]["thieu"] == 1 and by_d[16]["tong_kg"] == n1.tong_kg


@pytest.mark.asyncio
async def test_mcp_civil_and_ma_dinh_muc(tmp_path, kho_dir, sample_dxf):
    from fastmcp import Client
    from cadtroly.mcp.server import build_server
    tn, ct = make_cross_sections(tmp_path / "tn.dxf"), make_rebar_notes(tmp_path / "ct.dxf")
    async with Client(build_server()) as c:
        r = _data(await c.call_tool("takeoff_earthwork", {"path": str(tn), "out": str(tmp_path / "dd.xlsx")}))
        assert r["excel"] is None and r["can_hoi"] and r["huong_dan"]         # còn câu hỏi → không xuất
        r = _data(await c.call_tool("takeoff_earthwork", {
            "path": str(tn), "out": str(tmp_path / "dd.xlsx"), "ma_dinh_muc": {"dao": "MA-DAO-TEST"},
            "thieu_dien_tich": "bang_0"}))
        assert r["totals_m3"]["Đào nền"] == 690.0
        wb = load_workbook(r["excel"])
        assert wb.sheetnames == ["Tong hop", "Khoi luong doan", "Trac ngang", "Kiem soat"]
        codes = [row[1] for row in wb["Tong hop"].iter_rows(min_row=6, values_only=True)]
        assert "MA-DAO-TEST" in codes
        assert any("38/2026" in str(v) for row in wb["Kiem soat"].iter_rows(values_only=True) for v in row)

        r = _data(await c.call_tool("takeoff_rebar", {"path": str(ct), "out": str(tmp_path / "thep.xlsx"),
                                                      "don_vi_L": "mm", "trung_lap": "mot_lan"}))
        assert r["can_hoi"] == []
        assert r["theo_duong_kinh"]["Ø22"]["tong_kg"] > 0
        assert load_workbook(r["excel"]).sheetnames == ["Thong ke thep", "Tong hop", "Kiem soat"]

        P = str(sample_dxf)
        await c.call_tool("takeoff_propose", {"path": P})
        await c.call_tool("takeoff_decide", {"path": P, "kind": "block", "name": "WD_2", "cong_tac": "Cửa sổ",
                                             "don_vi": "bộ", "cach_tinh": "count", "ma_dinh_muc": "MA-CS-TEST"})
        rows = _data(await c.call_tool("takeoff_compute", {"path": P}))["rows"]
        assert any(x["ma_dinh_muc"] == "MA-CS-TEST" for x in rows)
    saved = json.loads((kho_dir / "quy-tac" / "anh-xa.json").read_text(encoding="utf-8"))
    assert any(r["ma_dinh_muc"] == "MA-CS-TEST" for r in saved)
